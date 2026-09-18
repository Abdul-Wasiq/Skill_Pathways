"""
Extracts raw text from uploaded resume files (PDF, DOCX).
This is purely mechanical text extraction — no AI involved here.
The extracted text is later sent to the AI service to be structured.
"""
import io

from docx import Document
from pypdf import PdfReader


class ResumeExtractionError(Exception):
    pass


def extract_text_from_pdf(file_bytes: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        pages_text = [page.extract_text() or "" for page in reader.pages]
        text = "\n".join(pages_text).strip()
    except Exception as exc:
        raise ResumeExtractionError(f"Could not read PDF file: {exc}") from exc
    if not text:
        raise ResumeExtractionError(
            "No extractable text found in this PDF. It may be a scanned image "
            "without OCR — try a text-based PDF or a DOCX file instead."
        )
    return text


def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also pull text out of tables, since resumes often use them for layout.
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text)
        text = "\n".join(paragraphs).strip()
    except Exception as exc:
        raise ResumeExtractionError(f"Could not read DOCX file: {exc}") from exc
    if not text:
        raise ResumeExtractionError("No extractable text found in this document.")
    return text


def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    if lower.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    raise ResumeExtractionError("Unsupported file type. Please upload a PDF or DOCX file.")
